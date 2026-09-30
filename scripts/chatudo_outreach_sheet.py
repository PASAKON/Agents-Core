#!/usr/bin/env python3
"""Build the Chatudo outreach tracker (plan work package O1) as an .xlsx that Drive converts to a Google Sheet.

Why an .xlsx: the Drive OAuth this org holds has the `drive` scope but the Sheets API is not enabled in its
Cloud project (probed 2026-09-28: 403 "Google Sheets API has not been used in project 407008779435"). A
Drive upload with mimeType application/vnd.google-apps.spreadsheet converts the file, formulas included.

Tabs (Thai labels, CEO reads them):
  ร้าน           one row per shop: segment, owner, has-LINE-OA, message variant, the draft, contact dates, day-3 /
                 day-7 follow-up due dates, the follow-up draft, replied, reached owner, demo, pilot, and a
                 "next action" formula.
  สรุปวันจันทร์    funnel counts for the Monday 10:00 review: per segment, per message variant, segment x
                 variant, per week, today's to-do counts, and the G1 check (31 Oct).
  ข้อความ         12 first messages keyed กลุ่ม|แบบ (4 segments x A/B/C) plus the day-3 and day-7 follow-ups, all
                 written by the CMO; {เจ้าของ} and {ร้าน} are filled in on the shop tab.
  วิธีใช้          how to use it.

Every count is a formula, so nothing has to be run on Monday. The drafts are a mail-merge formula ($0, no
model call).

Columns are addressed by name (COL["pilot"]), never by letter, so a new column only needs a row in SHOP_COLS.
Changed 2026-09-30 at the CMO's request (cmo-3c1fc6bc): a "มี LINE OA" column (November installs are LINE
only), 12 messages instead of 3, and a follow-up draft column.

Usage:
  python3 scripts/chatudo_outreach_sheet.py --out /path/tracker.xlsx            # empty tracker
  python3 scripts/chatudo_outreach_sheet.py --out /path/t.xlsx --sample 12      # local self-test rows
  python3 scripts/chatudo_outreach_sheet.py --out /path/t.xlsx --messages docs/promo/CHATUDO-MESSAGES-2026-10.txt
Needs openpyxl.
"""
import argparse
import datetime as dt
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROWS = 500  # formula rows on the shop tab (the plan lists ~450 shops)
FIRST = 2   # first data row on the shop tab
LAST = FIRST + ROWS - 1

SEGMENTS = [
    "คอร์สออนไลน์ / สัมมนา",
    "ร้านจองคิว (ทำผม เล็บ สปา)",
    "ร้านเล็กขายของในแชท",
    "คลินิกความงาม",
]
VARIANTS = ["A", "B", "C"]
YES, NO = "ใช่", "ไม่"
OWNER_SOURCES = ["หน้าความโปร่งใสของเพจ", "เว็บไซต์", "กรมพัฒนาธุรกิจการค้า", "หน้าแบรนด์ / ผู้สอน", "อื่น ๆ"]
CHANNELS = ["Inbox เพจ (FB ส่วนตัว CEO)", "อีเมล", "Inbox เพจ + อีเมล"]
FIRST_MONDAY = dt.date(2026, 9, 28)  # week 1 of the plan: outreach starts Thu 1 Oct
WEEKS = 10
G1_DATE = dt.date(2026, 10, 31)
FOLLOW1, FOLLOW2 = "ตามครั้งที่ 1", "ตามครั้งที่ 2"  # next-action labels; the follow-up draft keys on them

# (name, header, width, kind) in sheet order; kind drives validation / number format.
SHOP_COLS = [
    ("num", "#", 5, "int"),
    ("seg", "กลุ่ม", 24, "segment"),
    ("shop", "ชื่อร้าน / เพจ", 26, "text"),
    ("fb", "ลิงก์เพจ Facebook", 28, "text"),
    ("adlib", "ลิงก์ Meta Ad Library", 24, "text"),
    ("web", "เว็บไซต์ / LINE OA", 22, "text"),
    ("lineoa", "มี LINE OA (พ.ย. ติดตั้งได้แค่ LINE)", 13, "yesno"),
    ("owner", "ชื่อเจ้าของ", 18, "text"),
    ("owner_src", "ได้ชื่อเจ้าของจาก", 20, "owner_src"),
    ("email", "อีเมลเจ้าของ", 22, "text"),
    ("variant", "ข้อความแบบ", 10, "variant"),
    ("draft", "ร่างข้อความแรก (เติมชื่อให้แล้ว)", 48, "formula"),
    ("sent", "วันที่ทัก", 12, "date"),
    ("channel", "ช่องทาง", 22, "channel"),
    ("f1_due", "ตามครั้งที่ 1 ครบ (วันที่ 3)", 14, "formula_date"),
    ("f1_sent", "ตามครั้งที่ 1 ส่งแล้ว", 13, "date"),
    ("f2_due", "ตามครั้งที่ 2 ครบ (วันที่ 7)", 14, "formula_date"),
    ("f2_sent", "ตามครั้งที่ 2 ส่งแล้ว", 13, "date"),
    ("replied", "ร้านตอบ", 9, "yesno"),
    ("replied_at", "วันที่ตอบ", 12, "date"),
    ("reached", "ถึงเจ้าของ", 10, "yesno"),
    ("called", "โทรแล้ว (วันที่)", 13, "date"),
    ("demo", "เปิด demo (พักไว้)", 11, "yesno"),
    ("pilot", "ร้านนำร่อง", 10, "yesno"),
    ("installed", "วันติดตั้ง", 12, "date"),
    ("next", "ต้องทำต่อ", 18, "formula"),
    ("follow_draft", "ร่างข้อความตาม (เติมชื่อให้แล้ว)", 40, "formula"),
    ("week", "สัปดาห์ที่ทัก (จันทร์)", 13, "formula_date"),
    ("note", "หมายเหตุ", 30, "text"),
]
COL = {name: get_column_letter(i + 1) for i, (name, *_rest) in enumerate(SHOP_COLS)}
COLS = {COL[name]: (head, width, kind) for name, head, width, kind in SHOP_COLS}  # letter -> (header, width, kind)
LAST_COL = COL[SHOP_COLS[-1][0]]

DATE_FMT = "dd/mm/yyyy"
# COUNTIFS criteria. Not "<>": the local evaluator used for the self-test reads it differently from
# Sheets/Excel, and these two mean the same thing in all three (text of >=1 char; a real date).
NONEMPTY_TEXT = '"?*"'
HAS_DATE = '">0"'

HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(bold=True, color="FFFFFF")
CALC_FILL = PatternFill("solid", fgColor="F3F4F6")  # formula columns: do not type here
AMBER = PatternFill("solid", fgColor="FDE7C2")
GREEN = PatternFill("solid", fgColor="D6F2E0")
GREY = PatternFill("solid", fgColor="E5E7EB")

SHOP = "ร้าน"
SUM = "สรุปวันจันทร์"
MSG = "ข้อความ"
HOW = "วิธีใช้"

# Message tab layout: header on row 2, the 12 keyed first messages on rows 3-14, follow-ups below.
MSG_HEAD = 2
MSG_FIRST = 3
MSG_LAST = MSG_FIRST + len(SEGMENTS) * len(VARIANTS) - 1  # 14
MSG_F1 = MSG_LAST + 2  # 16
MSG_F2 = MSG_F1 + 1    # 17
MSG_KEY_SEP = "|"


def q(sheet):
    return f"'{sheet}'"


def rng(name):
    col = COL[name]
    return f"{q(SHOP)}!${col}${FIRST}:${col}${LAST}"


def fill_names(text_expr, r):
    """SUBSTITUTE {เจ้าของ} and {ร้าน} in a message expression for shop row r."""
    owner, shop = f"{COL['owner']}{r}", f"{COL['shop']}{r}"
    return (f'SUBSTITUTE(SUBSTITUTE({text_expr},"{{เจ้าของ}}",IF({owner}="","คุณเจ้าของร้าน","คุณ"&{owner})),'
            f'"{{ร้าน}}",{shop})')


def shop_formulas(r):
    c = lambda name: f"{COL[name]}{r}"  # noqa: E731
    key = f'{c("seg")}&"{MSG_KEY_SEP}"&{c("variant")}'
    lookup = f"VLOOKUP({key},{q(MSG)}!$C${MSG_FIRST}:$D${MSG_LAST},2,FALSE)"
    draft = (
        f'IF(OR({c("shop")}="",{c("seg")}="",{c("variant")}=""),"",'
        f'IFERROR(IF({lookup}="","",{fill_names(lookup, r)}),""))'
    )
    nxt = (
        f'IF({c("shop")}="","",'
        f'IF({c("pilot")}="{YES}",IF({c("installed")}="","ติดตั้งร้านนำร่อง","นำร่องแล้ว"),'
        f'IF({c("replied")}="{YES}",IF({c("called")}="","โทรหาเจ้าของ","รอผลหลังโทร"),'
        f'IF({c("sent")}="","ทักครั้งแรก",'
        f'IF(AND({c("f1_sent")}="",TODAY()>={c("f1_due")}),"{FOLLOW1}",'
        f'IF(AND({c("f1_sent")}<>"",{c("f2_sent")}="",TODAY()>={c("f2_due")}),"{FOLLOW2}",'
        f'IF({c("f2_sent")}<>"","หยุด (ตามครบ 2 ครั้ง)","รอ")))))))'
    )
    f1, f2 = f"{q(MSG)}!$D${MSG_F1}", f"{q(MSG)}!$D${MSG_F2}"
    follow = (
        f'IF({c("next")}="{FOLLOW1}",IF({f1}="","",{fill_names(f1, r)}),'
        f'IF({c("next")}="{FOLLOW2}",IF({f2}="","",{fill_names(f2, r)}),""))'
    )
    return {
        "draft": "=" + draft,
        "f1_due": f'=IF({c("sent")}="","",{c("sent")}+3)',
        "f2_due": f'=IF({c("sent")}="","",{c("sent")}+7)',
        "next": "=" + nxt,
        "follow_draft": "=" + follow,
        "week": f'=IF({c("sent")}="","",{c("sent")}-WEEKDAY({c("sent")},3))',
    }


def build_shop(ws, sample):
    for col, (head, width, kind) in COLS.items():
        cell = ws[f"{col}1"]
        cell.value = head
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{LAST_COL}{LAST}"

    for r in range(FIRST, LAST + 1):
        ws[f"{COL['num']}{r}"] = r - FIRST + 1
        for name, f in shop_formulas(r).items():
            ws[f"{COL[name]}{r}"] = f
        for col, (_, _, kind) in COLS.items():
            if kind.startswith("formula"):
                ws[f"{col}{r}"].fill = CALC_FILL
            if kind in ("date", "formula_date"):
                ws[f"{col}{r}"].number_format = DATE_FMT
        for name in ("draft", "follow_draft"):
            ws[f"{COL[name]}{r}"].alignment = Alignment(wrap_text=False)

    def dv(values, name):
        col = COL[name]
        v = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        v.add(f"{col}{FIRST}:{col}{LAST}")
        ws.add_data_validation(v)

    for col, (_, _, kind) in COLS.items():
        if kind == "date":
            d = DataValidation(type="date", operator="greaterThan", formula1="46023", allow_blank=True,
                               showErrorMessage=True, errorTitle="ใส่เป็นวันที่",
                               error="พิมพ์วันที่แบบ วว/ดด/ปปปป (ค.ศ.) เช่น 01/10/2026 หรือกด Ctrl + ; เพื่อใส่วันนี้")
            d.add(f"{col}{FIRST}:{col}{LAST}")
            ws.add_data_validation(d)
    dv(SEGMENTS, "seg")
    dv(OWNER_SOURCES, "owner_src")
    dv(VARIANTS, "variant")
    dv(CHANNELS, "channel")
    for name, _h, _w, kind in SHOP_COLS:
        if kind == "yesno":
            dv([YES, NO], name)

    nc = COL["next"]
    for text, fill in (("ทักครั้งแรก", AMBER), ("ตามครั้งที่", AMBER), ("โทรหา", GREEN),
                       ("ติดตั้ง", GREEN), ("หยุด", GREY)):
        ws.conditional_formatting.add(
            f"{nc}{FIRST}:{nc}{LAST}",
            FormulaRule(formula=[f'ISNUMBER(SEARCH("{text}",{nc}{FIRST}))'], fill=fill))

    if sample:
        fill_sample(ws, sample)


def fill_sample(ws, n):
    """Synthetic rows for the local self-test only; never left on Drive."""
    today = dt.date.today()

    def put(name, r, v):
        ws[f"{COL[name]}{r}"] = v

    for i in range(n):
        r = FIRST + i
        put("seg", r, SEGMENTS[i % 4])
        put("shop", r, f"ร้านทดสอบ {i + 1}")
        put("owner", r, "สมชาย" if i % 2 == 0 else None)
        put("variant", r, VARIANTS[i % 3])
        put("lineoa", r, YES if i % 3 != 2 else NO)
        if i < n - 2:  # the last two rows are not contacted yet
            put("sent", r, today - dt.timedelta(days=i))
        if i in (4, 5, 6, 7, 8, 9):
            put("f1_sent", r, today - dt.timedelta(days=i - 3))
        if i in (8, 9):
            put("f2_sent", r, today)
        if i in (1, 2):
            put("replied", r, YES)
            put("replied_at", r, today)
            put("reached", r, YES)
        if i == 2:
            put("called", r, today)
            put("pilot", r, YES)


def parse_messages(path):
    """Read the CMO's message file: '### <A|B|C> / <segment>' blocks plus '### ตามครั้งที่ 1…' / '### ตามครั้งที่ 2…'.

    Returns ({(segment, variant): text}, follow1, follow2). Refuses a missing block, an unknown segment or an em
    dash (IRON §39: no em dash in Thai public copy).
    """
    blocks, head, lines = [], None, []
    for raw in open(path, encoding="utf-8").read().splitlines():
        if raw.startswith("### "):
            if head is not None:
                blocks.append((head, "\n".join(lines).strip()))
            head, lines = raw[4:].strip(), []
        elif head is not None:
            lines.append(raw.rstrip())
    if head is not None:
        blocks.append((head, "\n".join(lines).strip()))

    firsts, follow = {}, {}
    for head, text in blocks:
        if "—" in text:
            raise SystemExit(f"em dash in message {head!r}: IRON §39 forbids it in Thai public copy")
        m = re.match(r"^([ABC])\s*/\s*(.+)$", head)
        if m:
            variant, seg = m.group(1), m.group(2).strip()
            if seg not in SEGMENTS:
                raise SystemExit(f"unknown segment {seg!r} in {head!r}; expected one of {SEGMENTS}")
            firsts[(seg, variant)] = text
        elif head.startswith(FOLLOW1):
            follow[1] = text
        elif head.startswith(FOLLOW2):
            follow[2] = text
    missing = [f"{v} / {s}" for s in SEGMENTS for v in VARIANTS if not firsts.get((s, v))]
    missing += [f"{lab}" for k, lab in ((1, FOLLOW1), (2, FOLLOW2)) if not follow.get(k)]
    if missing:
        raise SystemExit(f"{path}: missing or empty blocks: {missing}")
    return firsts, follow[1], follow[2]


def build_messages(ws, sample=0, messages=None):
    ws["A1"] = ("ข้อความแรก 12 ชุด (4 กลุ่ม x แบบ A/B/C) กับข้อความตามวันที่ 3 และ 7. CMO เขียน, CEO ส่ง. "
                "ใส่ {เจ้าของ} กับ {ร้าน} ตรงที่ให้เติมชื่อ: {เจ้าของ} จะเป็น 'คุณ<ชื่อเจ้าของ>' หรือ "
                "'คุณเจ้าของร้าน' ถ้ายังไม่รู้ชื่อ. แก้ข้อความได้ในคอลัมน์ D; อย่าแก้คอลัมน์ A-C")
    ws["A1"].font = Font(bold=True)
    for col, head in zip("ABCDEF", ("กลุ่ม", "แบบ", "key (คำนวณเอง)", "ข้อความ", "ผู้เขียน", "แก้ล่าสุด")):
        ws[f"{col}{MSG_HEAD}"] = head
        ws[f"{col}{MSG_HEAD}"].fill, ws[f"{col}{MSG_HEAD}"].font = HEAD_FILL, HEAD_FONT
    firsts, f1, f2 = messages if messages else ({}, None, None)
    r = MSG_FIRST
    for seg in SEGMENTS:
        for v in VARIANTS:
            ws[f"A{r}"], ws[f"B{r}"] = seg, v
            ws[f"C{r}"] = f'=A{r}&"{MSG_KEY_SEP}"&B{r}'
            ws[f"C{r}"].fill = CALC_FILL
            # CMO fills; an empty row leaves that draft empty. Sample text is for the self-test only.
            if sample:
                ws[f"D{r}"] = f"[ทดสอบ {seg}{MSG_KEY_SEP}{v}] สวัสดีครับ {{เจ้าของ}} ร้าน {{ร้าน}}"
            else:
                ws[f"D{r}"] = firsts.get((seg, v))
            ws[f"E{r}"] = "CMO"
            ws[f"D{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    ws[f"A{MSG_F1}"] = f"{FOLLOW1} (วันที่ 3)"
    ws[f"A{MSG_F2}"] = f"{FOLLOW2} (วันที่ 7, ครั้งสุดท้าย)"
    for row, text, tag in ((MSG_F1, f1, "1"), (MSG_F2, f2, "2")):
        ws[f"B{row}"] = "ทุกกลุ่ม"
        ws[f"D{row}"] = f"[ทดสอบ ตาม {tag}] {{เจ้าของ}} ร้าน {{ร้าน}}" if sample else text
        ws[f"E{row}"] = "CMO"
        ws[f"D{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 90
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 12


def cnt(*pairs):
    """COUNTIFS over the shop tab; pairs of (column name, criterion-expression)."""
    args = ",".join(f"{rng(name)},{crit}" for name, crit in pairs)
    return f"COUNTIFS({args})"


def pct(num, den):
    return f'=IF({den}=0,"",{num}/{den})'


def build_summary(ws):
    # Letters below are cells of THIS summary tab; shop-tab columns go through cnt() by name.
    bold = Font(bold=True)
    ws["A1"] = "สรุปสำหรับทบทวนทุกวันจันทร์ 10:00 (คำนวณเองทั้งหมด ไม่ต้องกรอก)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "วันนี้"
    ws["B2"] = "=TODAY()"
    ws["B2"].number_format = DATE_FMT

    def header(row, labels):
        for j, lab in enumerate(labels):
            cell = ws.cell(row=row, column=1 + j, value=lab)
            cell.fill, cell.font = HEAD_FILL, HEAD_FONT
            cell.alignment = Alignment(wrap_text=True)

    # 1. per segment, cumulative
    r0 = 4
    ws[f"A{r0}"] = "1. สะสมตามกลุ่ม"
    ws[f"A{r0}"].font = bold
    header(r0 + 1, ["กลุ่ม", "ในรายชื่อ", "ทักแล้ว", "ร้านตอบ", "% ตอบ", "ถึงเจ้าของ", "% ถึงเจ้าของ",
                    "เปิด demo", "ร้านนำร่อง", "ติดตั้งแล้ว", "มี LINE OA"])
    rows = SEGMENTS + ["รวม"]
    for i, seg in enumerate(rows):
        r = r0 + 2 + i
        ws[f"A{r}"] = seg
        if seg == "รวม":
            first, last = r0 + 2, r - 1
            for col in "BCDFHIJK":
                ws[f"{col}{r}"] = f"=SUM({col}{first}:{col}{last})"
            ws[f"A{r}"].font = bold
        else:
            s = f"$A{r}"
            ws[f"B{r}"] = "=" + cnt(("seg", s), ("shop", NONEMPTY_TEXT))
            ws[f"C{r}"] = "=" + cnt(("seg", s), ("sent", HAS_DATE))
            ws[f"D{r}"] = "=" + cnt(("seg", s), ("replied", f'"{YES}"'))
            ws[f"F{r}"] = "=" + cnt(("seg", s), ("reached", f'"{YES}"'))
            ws[f"H{r}"] = "=" + cnt(("seg", s), ("demo", f'"{YES}"'))
            ws[f"I{r}"] = "=" + cnt(("seg", s), ("pilot", f'"{YES}"'))
            ws[f"J{r}"] = "=" + cnt(("seg", s), ("pilot", f'"{YES}"'), ("installed", HAS_DATE))
            ws[f"K{r}"] = "=" + cnt(("seg", s), ("lineoa", f'"{YES}"'))
        ws[f"E{r}"] = pct(f"D{r}", f"C{r}")
        ws[f"G{r}"] = pct(f"F{r}", f"C{r}")
        for col in "EG":
            ws[f"{col}{r}"].number_format = "0%"
    seg_total_row = r0 + 2 + len(SEGMENTS)

    # 2. per message variant
    r1 = seg_total_row + 3
    ws[f"A{r1}"] = "2. ตามแบบข้อความ (ทดสอบ 3 แบบ)"
    ws[f"A{r1}"].font = bold
    header(r1 + 1, ["แบบ", "ทักแล้ว", "ร้านตอบ", "% ตอบ", "ถึงเจ้าของ", "% ถึงเจ้าของ", "ร้านนำร่อง"])
    for i, v in enumerate(VARIANTS):
        r = r1 + 2 + i
        ws[f"A{r}"] = v
        s = f"$A{r}"
        ws[f"B{r}"] = "=" + cnt(("variant", s), ("sent", HAS_DATE))
        ws[f"C{r}"] = "=" + cnt(("variant", s), ("replied", f'"{YES}"'))
        ws[f"D{r}"] = pct(f"C{r}", f"B{r}")
        ws[f"E{r}"] = "=" + cnt(("variant", s), ("reached", f'"{YES}"'))
        ws[f"F{r}"] = pct(f"E{r}", f"B{r}")
        ws[f"G{r}"] = "=" + cnt(("variant", s), ("pilot", f'"{YES}"'))
        for col in "DF":
            ws[f"{col}{r}"].number_format = "0%"

    # 3. segment x variant: % reached owner
    r2 = r1 + 2 + len(VARIANTS) + 2
    ws[f"A{r2}"] = "3. % ถึงเจ้าของ แยกกลุ่ม x แบบข้อความ (ถึงเจ้าของ / ทักแล้ว)"
    ws[f"A{r2}"].font = bold
    header(r2 + 1, ["กลุ่ม"] + [f"แบบ {v}" for v in VARIANTS])
    for i, seg in enumerate(SEGMENTS):
        r = r2 + 2 + i
        ws[f"A{r}"] = seg
        for j, v in enumerate(VARIANTS):
            col = get_column_letter(2 + j)
            sent = cnt(("seg", f"$A{r}"), ("variant", f'"{v}"'), ("sent", HAS_DATE))
            hit = cnt(("seg", f"$A{r}"), ("variant", f'"{v}"'), ("reached", f'"{YES}"'))
            ws[f"{col}{r}"] = f'=IF({sent}=0,"",TEXT({hit}/{sent},"0%")&" ("&{hit}&"/"&{sent}&")")'

    # 4. per week (Monday to Sunday)
    r3 = r2 + 2 + len(SEGMENTS) + 2
    ws[f"A{r3}"] = "4. รายสัปดาห์ (จันทร์ถึงอาทิตย์)"
    ws[f"A{r3}"].font = bold
    header(r3 + 1, ["สัปดาห์เริ่ม", "ทักใหม่", "ร้านตอบ (ตามวันที่ตอบ)", "ถึงเจ้าของ (ตามสัปดาห์ที่ทัก)",
                    "ติดตั้งร้านนำร่อง (ตามวันติดตั้ง)", "เป้าทัก (20-30/วัน)"])
    for i in range(WEEKS):
        r = r3 + 2 + i
        ws[f"A{r}"] = FIRST_MONDAY + dt.timedelta(weeks=i)
        ws[f"A{r}"].number_format = DATE_FMT
        wk = f"$A{r}"
        ws[f"B{r}"] = "=" + cnt(("week", wk))
        ws[f"C{r}"] = "=" + cnt(("replied_at", f'">="&{wk}'), ("replied_at", f'"<"&{wk}+7'))
        ws[f"D{r}"] = "=" + cnt(("week", wk), ("reached", f'"{YES}"'))
        ws[f"E{r}"] = "=" + cnt(("installed", f'">="&{wk}'), ("installed", f'"<"&{wk}+7'))
        ws[f"F{r}"] = "80-120" if i == 0 else "140-210"  # week 1 starts Thu 1 Oct: 4 days

    # 5. today
    r4 = r3 + 2 + WEEKS + 1
    ws[f"A{r4}"] = "5. ต้องทำวันนี้ (กรองคอลัมน์ 'ต้องทำต่อ' ในแท็บร้าน)"
    ws[f"A{r4}"].font = bold
    for i, label in enumerate(["ทักครั้งแรก", FOLLOW1, FOLLOW2, "โทรหาเจ้าของ", "ติดตั้งร้านนำร่อง"]):
        r = r4 + 1 + i
        ws[f"A{r}"] = label
        ws[f"B{r}"] = "=" + cnt(("next", f'"{label}"'))

    # 6. G1
    r5 = r4 + 7
    ws[f"A{r5}"] = f"6. ประตู G1 ({G1_DATE:%d/%m/%Y} 18:00) ผ่านเมื่อครบทุกข้อ"
    ws[f"A{r5}"].font = bold
    header(r5 + 1, ["เกณฑ์", "ตอนนี้", "เป้า", "ผ่าน"])
    tot = seg_total_row
    g1 = [
        ("ทักแล้ว (ร้าน)", f"=C{tot}", 400, None),
        ("% ถึงเจ้าของ", f"=IF(C{tot}=0,0,F{tot}/C{tot})", 0.10, "0%"),
        ("ร้านนำร่องติดตั้ง", f"=J{tot}", 7, None),
    ]
    for i, (lab, val, goal, fmt) in enumerate(g1):
        r = r5 + 2 + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = lab, val, goal
        ws[f"D{r}"] = f'=IF(B{r}>=C{r},"{YES}","ยัง")'
        if fmt:
            ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = fmt
    ws[f"A{r5 + 5}"] = "เลือก 2 กลุ่มที่ดีที่สุดจากตาราง 1 (% ถึงเจ้าของ และจำนวนร้านนำร่อง)"

    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 14
    return {"seg_total_row": seg_total_row, "variant_row": r1 + 2, "week_row": r3 + 2,
            "today_row": r4 + 1, "g1_row": r5 + 2}


def build_howto(ws):
    lines = [
        "วิธีใช้ตารางทักร้าน Chatudo",
        "1. แท็บ 'ร้าน': กรอกหนึ่งแถวต่อร้าน (กลุ่ม, ชื่อร้าน, ลิงก์, มี LINE OA, ชื่อเจ้าของ, แบบข้อความ A/B/C). "
        "ช่องสีเทาคำนวณเอง ไม่ต้องพิมพ์",
        "2. คัดลอก 'ร่างข้อความแรก' ไปส่ง (เลือกตามกลุ่ม + แบบ) แล้วใส่ 'วันที่ทัก'. วันตามครั้งที่ 1 (วันที่ 3) และ"
        "ครั้งที่ 2 (วันที่ 7) ขึ้นเอง",
        "3. ทุกเช้า กรองคอลัมน์ 'ต้องทำต่อ': ทักครั้งแรก / ตามครั้งที่ 1 / ตามครั้งที่ 2 / โทรหาเจ้าของ. "
        "ข้อความตามอยู่ในคอลัมน์ 'ร่างข้อความตาม'. ส่งแล้วใส่วันที่ในช่อง 'ส่งแล้ว'",
        "4. ร้านตอบ: ใส่ 'ใช่' + วันที่ตอบ. ถึงตัวเจ้าของจริง: ช่อง 'ถึงเจ้าของ' = ใช่. โทรเฉพาะร้านที่ตอบแล้ว. "
        "ตามครบ 2 ครั้งแล้วไม่ตอบ = หยุด",
        "5. 'มี LINE OA' = ใช่ เฉพาะร้านที่มี LINE OA แล้ว. เดือน พ.ย. ติดตั้งได้แค่ร้านที่มี LINE OA",
        "6. วันจันทร์ 10:00 เปิดแท็บ 'สรุปวันจันทร์' ตัวเลขทุกช่องนับให้แล้ว. ข้อความ 12 ชุดกับข้อความตามอยู่ในแท็บ "
        "'ข้อความ' (CMO เขียน)",
        "ที่มา: แผน Chatudo งาน O1 (CTO), CMO แผน #plan. ห้ามแทรกหรือลบคอลัมน์ในแท็บ 'ร้าน' เพราะสูตรสรุปอ้างตามตำแหน่ง",
    ]
    for i, t in enumerate(lines):
        ws[f"A{1 + i}"] = t
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 140


def build(out, sample=0, messages_path=None):
    messages = parse_messages(messages_path) if messages_path else None
    wb = Workbook()
    ws = wb.active
    ws.title = SHOP
    build_shop(ws, sample)
    layout = build_summary(wb.create_sheet(SUM))
    build_messages(wb.create_sheet(MSG), sample, messages)
    build_howto(wb.create_sheet(HOW))
    wb.save(out)
    return layout


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--sample", type=int, default=0, help="synthetic rows for a local self-test only")
    ap.add_argument("--messages", help="the CMO's message file (### <A|B|C> / <segment> blocks + 2 follow-ups)")
    a = ap.parse_args()
    if a.sample and a.messages:
        ap.error("--sample and --messages are exclusive: sample rows use test text")
    print(build(a.out, a.sample, a.messages))
