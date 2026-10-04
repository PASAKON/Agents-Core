#!/usr/bin/env python3
"""Put the Chatudo outreach tracker (plan O1) on Drive as a Google Sheet, and check what Google computes.

  verify FILE_ID   read-only: export the live Sheet, count formula errors, compare its TODAY() with Bangkok's date
  add-shops FILE_ID CSV [--apply]
                   preview free shop rows; --apply appends only mapped CSV cells using RAW values,
                   checks target cells again before writing, then verifies input and formula errors.
  run              ensure PROJECT/CHATUDO/Sales & Outreach (CEO approved 2026-09-28), create the Sheet from a
                   12-row sample build, check Google's values, then replace the content with the empty tracker
                   (same file id, so nothing is deleted) and check again. Refuses if the Sheet already exists.
  update FILE_ID [--messages FILE]
                   rebuild an existing Sheet in place after a layout change (same file id and link; Drive keeps
                   the old version in its history). Refuses if any shop row holds input or a message cell was
                   typed on Drive, so it never overwrites the CEO's work. Loads the sample build, checks Google's
                   values, then loads the empty build (with the CMO's messages when --messages is given).

Live Sheet (2026-09-28): 1kqCYwXOJ2rdnOwAtitICTtoZ3WW01aBuz1cMP6Ta7e0 in PROJECT/CHATUDO/Sales & Outreach.

A converted Sheet starts on America/Los_Angeles, so TODAY() lags Bangkok until 14:00 and day-3/day-7 follow-ups
can show a day late. CEO 2026-09-28: do not chase the Sheet's time zone or locale. Any agent that summarises the
Sheet reports in Thai time (Asia/Bangkok) and works out "today" itself. `verify` still reports the clock, but it
can tell only between 00:00 and 14:00 Bangkok, when the two dates differ; outside that window it says INCONCLUSIVE.

Auth: the Drive OAuth in ClaudeFlow's .env, through scripts/gdrive-bridge/ilag_sync.py (values never printed).
Needs openpyxl. Work files go to $CHATUDO_O1_DIR (default: the system temp dir).
"""
import csv
import datetime as dt
import json
import os
import subprocess
import sys
import tempfile
import urllib.parse
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.environ.get("CHATUDO_O1_DIR") or tempfile.gettempdir()
os.environ.setdefault("MOONIEX_CLAUDEFLOW_ENV", "/opt/MoonieXHQ/Projects/MoonieX/ClaudeFlow/.env")
sys.path.insert(0, os.path.join(HERE, "gdrive-bridge"))
import ilag_sync as g  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

sys.path.insert(0, HERE)
import chatudo_outreach_sheet as sheet  # noqa: E402

PROJECT_ID = "1HVLTPS09R4WvhYyOPW_0HHrfWxWGF-hQ"
L1, L2 = "CHATUDO", "Sales & Outreach"
TITLE = "Chatudo ตารางทักร้าน"
FOLDER = "application/vnd.google-apps.folder"
GSHEET = "application/vnd.google-apps.spreadsheet"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ERRS = {"#NAME?", "#ERROR!", "#VALUE!", "#REF!", "#N/A", "#DIV/0!", "#NUM!", "#NULL!"}
SAMPLE_N = 12


def expected_sample(build_day, sheet_today, n=SAMPLE_N):
    """What Google should compute for chatudo_outreach_sheet.fill_sample(n), worked out here independently.

    The converted Sheet runs on Pacific time, so its TODAY() (sheet_today) can be a day behind the build day
    before 14:00 Bangkok; the expectation follows the Sheet's own clock rather than a hard-coded list.
    """
    seg, var, col = sheet.SEGMENTS, sheet.VARIANTS, sheet.COL
    out = {}
    for i in range(n):
        r = sheet.FIRST + i
        sent = build_day - dt.timedelta(days=i) if i < n - 2 else None
        f1_sent = build_day - dt.timedelta(days=i - 3) if i in (4, 5, 6, 7, 8, 9) else None
        f2_sent = build_day if i in (8, 9) else None
        replied, called, pilot = i in (1, 2), i == 2, i == 2
        if pilot:
            nxt = "ติดตั้งร้านนำร่อง"
        elif replied:
            nxt = "รอผลหลังโทร" if called else "โทรหาเจ้าของ"
        elif sent is None:
            nxt = "ทักครั้งแรก"
        elif f1_sent is None and sheet_today >= sent + dt.timedelta(days=3):
            nxt = sheet.FOLLOW1
        elif f1_sent is not None and f2_sent is None and sheet_today >= sent + dt.timedelta(days=7):
            nxt = sheet.FOLLOW2
        elif f2_sent is not None:
            nxt = "หยุด (ตามครบ 2 ครั้ง)"
        else:
            nxt = "รอ"
        owner = "คุณสมชาย" if i % 2 == 0 else "คุณเจ้าของร้าน"
        shop = f"ร้านทดสอบ {i + 1}"
        out[f"{col['next']}{r}"] = nxt
        out[f"{col['draft']}{r}"] = f"[ทดสอบ {seg[i % 4]}{sheet.MSG_KEY_SEP}{var[i % 3]}] สวัสดีครับ {owner} ร้าน {shop}"
        tag = {sheet.FOLLOW1: "1", sheet.FOLLOW2: "2"}.get(nxt)
        out[f"{col['follow_draft']}{r}"] = f"[ทดสอบ ตาม {tag}] {owner} ร้าน {shop}" if tag else None
    return out


def check_sample(wb, build_day):
    """Mismatches between Google's computed sample values and expected_sample(); [] means the formulas hold."""
    sheet_today = day(wb[sheet.SUM]["B2"].value)
    want = expected_sample(build_day, dt.date.fromisoformat(sheet_today))
    shop = wb[sheet.SHOP]
    got = {k: (shop[k].value if shop[k].value != "" else None) for k in want}
    fails = [f"{k}: {got[k]!r} != {v!r}" for k, v in want.items() if got[k] != v]
    lineoa_total = wb[sheet.SUM][f"K{4 + 2 + len(sheet.SEGMENTS)}"].value
    want_lineoa = sum(1 for i in range(SAMPLE_N) if i % 3 != 2)
    if lineoa_total != want_lineoa:
        fails.append(f"summary มี LINE OA total {lineoa_total!r} != {want_lineoa}")
    return fails



def children(fid):
    res = g.api(g.DRIVE_FILES, params={"q": f"'{fid}' in parents and trashed = false",
                                       "fields": "files(id,name,mimeType,webViewLink)", "pageSize": "200"})
    return res.get("files", [])


def ensure_folder(parent, name):
    for f in children(parent):
        if f["name"] == name and f["mimeType"] == FOLDER:
            return f["id"], False
    made = g.api(g.DRIVE_FILES, method="POST", params={"fields": "id"},
                 data=json.dumps({"name": name, "parents": [parent], "mimeType": FOLDER}).encode(),
                 headers={"Content-Type": "application/json"})
    return made["id"], True


def build(path, sample, messages=None):
    cmd = [sys.executable, os.path.join(HERE, "chatudo_outreach_sheet.py"), "--out", path]
    if sample:
        cmd += ["--sample", str(sample)]
    if messages:
        cmd += ["--messages", messages]
    subprocess.run(cmd, check=True)


def create_sheet(path, parent):
    boundary = "o1chatudoboundary"
    meta = json.dumps({"name": TITLE, "parents": [parent], "mimeType": GSHEET}).encode()
    body = (f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n".encode() + meta
            + f"\r\n--{boundary}\r\nContent-Type: {XLSX}\r\n\r\n".encode() + open(path, "rb").read()
            + f"\r\n--{boundary}--".encode())
    return g.api(g.DRIVE_UPLOAD, method="POST",
                 params={"uploadType": "multipart", "fields": "id,name,mimeType,webViewLink"},
                 data=body, headers={"Content-Type": f"multipart/related; boundary={boundary}"})


def replace_content(fid, path):
    return g.api(f"{g.DRIVE_UPLOAD}/{fid}", method="PATCH",
                 params={"uploadType": "media", "fields": "id,name,mimeType,modifiedTime"},
                 data=open(path, "rb").read(), headers={"Content-Type": XLSX})


def export(fid, out):
    url = f"{g.DRIVE_FILES}/{fid}/export?" + urllib.parse.urlencode({"mimeType": XLSX})
    req = urllib.request.Request(url, headers={"Authorization": "Bearer " + g.access_token()})
    with urllib.request.urlopen(req, timeout=120) as r:
        open(out, "wb").write(r.read())
    return load_workbook(out, data_only=True)


def errors(wb):
    return [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if c.data_type == "e" or (isinstance(c.value, str) and c.value.strip() in ERRS)]


def day(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else (v.isoformat() if isinstance(v, dt.date) else v)


def main():
    step = sys.argv[1] if len(sys.argv) > 1 else ""
    if step == "add-shops":
        args = sys.argv[2:]
        apply = len(args) == 3 and args[-1] == "--apply"
        if len(args) != 2 and not apply:
            raise SystemExit("usage: add-shops FILE_ID CSV [--apply]")
        add_shops(args[0], args[1], apply=apply)
        return
    if step == "verify":
        wb = export(sys.argv[2], os.path.join(WORK, "o1-export-verify.xlsx"))
        now_bkk = dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=7)
        bkk = now_bkk.date().isoformat()
        sheet_today = day(wb["สรุปวันจันทร์"]["B2"].value)
        # From 14:00 Bangkok, Pacific time is on the same date, so a match proves nothing then.
        if now_bkk.hour >= 14:
            verdict = "INCONCLUSIVE (run between 00:00 and 14:00 Bangkok)"
        else:
            verdict = "TZ OK" if sheet_today == bkk else "TZ WRONG"
        print(f"errors={len(errors(wb))} sheet TODAY()={sheet_today} bangkok={bkk} "
              f"{now_bkk:%H:%M} -> {verdict}")
        return
    if step == "update":
        args = sys.argv[2:]
        msgs = None
        if "--messages" in args:
            k = args.index("--messages")
            msgs = args[k + 1]
            del args[k:k + 2]
        if len(args) != 1:
            raise SystemExit(__doc__)
        update(args[0], msgs)
        return
    if step != "run":
        raise SystemExit(__doc__)

    l1, made1 = ensure_folder(PROJECT_ID, L1)
    l2, made2 = ensure_folder(l1, L2)
    print(f"L1 {L1} {l1} {'CREATED' if made1 else 'existing'}")
    print(f"L2 {L2} {l2} {'CREATED' if made2 else 'existing'}")
    if any(f["name"] == TITLE for f in children(l2)):
        raise SystemExit(f"{TITLE!r} already exists in {L2}; use verify")

    sample, empty = os.path.join(WORK, "o1-sample.xlsx"), os.path.join(WORK, "o1-empty.xlsx")
    build_day = dt.date.today()
    build(sample, SAMPLE_N)
    build(empty, 0)
    f = create_sheet(sample, l2)
    fid = f["id"]
    wb = export(fid, os.path.join(WORK, "o1-export-sample.xlsx"))
    print(f"sample: errors={errors(wb)[:10]} mismatches={check_sample(wb, build_day)}")

    replace_content(fid, empty)
    wb2 = export(fid, os.path.join(WORK, "o1-export-empty.xlsx"))
    print(f"empty: errors={errors(wb2)[:10]} sample cells left={sample_left(wb2)} tabs={wb2.sheetnames}")
    print("SHEET", fid, f.get("webViewLink"))


def sample_left(wb):
    cols = [sheet.COL[k] for k in ("seg", "shop", "sent", "next")]
    return [f"{c}{r}" for r in range(sheet.FIRST, sheet.FIRST + SAMPLE_N) for c in cols
            if wb[sheet.SHOP][f"{c}{r}"].value not in (None, "")]


def user_input(wb):
    """Cells a person typed into: any non-formula shop column with a value, or a filled message text cell."""
    shop = wb[sheet.SHOP]
    formula_heads = {head for head, _w, kind in sheet.COLS.values() if kind.startswith("formula")}
    typed = [c for c in range(2, shop.max_column + 1)
             if shop.cell(row=1, column=c).value and shop.cell(row=1, column=c).value not in formula_heads]
    found = [f"{shop.cell(row=1, column=c).value}!{r}" for r in range(sheet.FIRST, shop.max_row + 1)
             for c in typed if shop.cell(row=r, column=c).value not in (None, "")]
    msg = wb[sheet.MSG]
    text_col = next((c for c in range(1, msg.max_column + 1) if msg.cell(row=2, column=c).value == "ข้อความ"), None)
    if text_col:
        found += [f"{sheet.MSG}!{msg.cell(row=r, column=text_col).coordinate}" for r in range(3, msg.max_row + 1)
                  if msg.cell(row=r, column=text_col).value not in (None, "")]
    return found


def add_shops(fid, path, apply=False):
    """Append CSV inputs to existing free rows; all diagnostics exclude cell values."""
    headers = {head: (sheet.COL[name], kind) for name, head, _width, kind in sheet.SHOP_COLS}
    with open(path, encoding="utf-8-sig", newline="") as source:
        reader = csv.reader(source, strict=True)
        heads = next(reader, [])
        if not heads or len(set(heads)) != len(heads):
            raise SystemExit("refusing: missing or duplicate CSV headers")
        if any(h not in headers or h == "#" or headers[h][1].startswith("formula") for h in heads):
            raise SystemExit("refusing: unknown or protected CSV header")
        rows = [[v.strip() for v in row] for row in reader]
    if not rows or any(len(row) != len(heads) for row in rows):
        raise SystemExit("refusing: empty CSV or inconsistent row width")
    cols = [headers[h][0] for h in heads]
    editable = [sheet.COL[name] for name, _h, _w, kind in sheet.SHOP_COLS
                if name != "num" and not kind.startswith("formula")]
    with tempfile.TemporaryDirectory(prefix="o1-add-shops-", dir=WORK) as work:
        before = export(fid, os.path.join(work, "before.xlsx"))
        shop = before[sheet.SHOP]
        targets = [r for r in range(sheet.FIRST, shop.max_row + 1)
                   if all(shop[f"{c}{r}"].value in (None, "") for c in editable)][:len(rows)]
        if len(targets) != len(rows):
            raise SystemExit("refusing: insufficient free rows")
        print(f"rows={len(rows)} target_rows={','.join(map(str, targets))} columns={','.join(cols)}")
        if not apply:
            return
        data = [{"range": f"'{sheet.SHOP}'!{col}{r}", "values": [[value]]}
                for r, row in zip(targets, rows) for col, value in zip(cols, row)]
        base = "https://sheets.googleapis.com/v4/spreadsheets/" + urllib.parse.quote(fid, safe="")
        query = urllib.parse.urlencode([("ranges", item["range"]) for item in data]
                                      + [("valueRenderOption", "FORMULA")])
        current = g.api(base + "/values:batchGet?" + query)
        ranges = current.get("valueRanges", [])
        if len(ranges) != len(data) or any(
                value not in (None, "") for item in ranges
                for row in item.get("values", []) for value in row):
            raise SystemExit("refusing: target cells changed or race check incomplete")
        g.api(base + "/values:batchUpdate", method="POST",
              data=json.dumps({"valueInputOption": "RAW", "data": data}).encode(),
              headers={"Content-Type": "application/json"})
        after = export(fid, os.path.join(work, "after.xlsx"))
        def normalized(value):
            return "" if value is None else value
        mismatches = sum(normalized(after[sheet.SHOP][f"{c}{r}"].value) != value
                         for r, row in zip(targets, rows) for c, value in zip(cols, row))
        prefix = sheet.SHOP + "!"
        new_errors = {e for e in errors(after) if e.startswith(prefix)} - {
            e for e in errors(before) if e.startswith(prefix)}
        # Build the expected input state, including old values, without changing the live Sheet.
        for r, row in zip(targets, rows):
            for c, value in zip(cols, row):
                shop[f"{c}{r}"] = value
        input_changed = set(user_input(after)) != set(user_input(before))
        def inputs(wb):
            return {(ws.title, cell.coordinate): cell.value for ws in wb.worksheets
                    for row in ws.iter_rows() for cell in row
                    if cell.value not in (None, "") and
                    (ws.title == sheet.MSG or (ws.title == sheet.SHOP and
                     cell.column_letter in editable and cell.row >= sheet.FIRST))}
        input_changed = input_changed or inputs(after) != inputs(before)
        print(f"verified_cells={len(data)} mismatches={mismatches} new_shop_errors={len(new_errors)} "
              f"input_changes={int(input_changed)}")
        if mismatches or new_errors or input_changed:
            raise SystemExit("post-write verification FAILED; writes were submitted")


def update(fid, messages=None):
    before = export(fid, os.path.join(WORK, f"o1-before-update-{dt.date.today():%Y%m%d}.xlsx"))
    typed = user_input(before)
    if typed:
        raise SystemExit(f"refusing: the live Sheet holds input that a rebuild would overwrite: {typed[:10]}")
    build_day = dt.date.today()
    sample, final = os.path.join(WORK, "o1-sample.xlsx"), os.path.join(WORK, "o1-final.xlsx")
    build(sample, SAMPLE_N)
    build(final, 0, messages)
    replace_content(fid, sample)
    wb = export(fid, os.path.join(WORK, "o1-export-sample.xlsx"))
    errs, fails = errors(wb), check_sample(wb, build_day)
    print(f"sample: errors={errs[:10]} mismatches={fails}")
    replace_content(fid, final)
    wb2 = export(fid, os.path.join(WORK, "o1-export-final.xlsx"))
    msg_filled = sum(1 for r in range(sheet.MSG_FIRST, sheet.MSG_F2 + 1) if wb2[sheet.MSG][f"D{r}"].value)
    print(f"final: errors={errors(wb2)[:10]} sample cells left={sample_left(wb2)} "
          f"message cells filled={msg_filled} tabs={wb2.sheetnames}")
    if errs or fails:
        raise SystemExit("sample check FAILED; the final build is loaded but its formulas are unproven")


if __name__ == "__main__":
    main()
